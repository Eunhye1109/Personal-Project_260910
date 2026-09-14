# -*- coding: utf-8 -*-
"""
발표자료에 두 장을 덧붙인다.

  · 분석 방식과 성능 지표   (결론 3 뒤)
  · 대시보드를 만든 이유     (대시보드 화면 소개 앞)

기존 슬라이드는 손대지 않는다. 새 장은 「한계와 후속 과제」 장을 통째로 복제해 만들기
때문에 글꼴·크기·색·여백이 기존 장과 정확히 같다. 글자만 갈아 끼운다.

실행
    python docs/add_method_slides.py
"""
from __future__ import annotations

import copy
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
DECK = ROOT / "docs" / "[BI 분석가 과정] 발표자료_260914_최은혜.pptx"

TEMPLATE = 15          # 0-based. 「한계와 후속 과제」 = 16쪽
DASHBOARD_TITLE = "대시보드: 화면 하나당 분석 질문 하나"

METHOD = {
    "eyebrow": "4. 결론 3 · 영상물",
    "title": "분석 방식과 성능 지표",
    "rows": [
        ["분석 단계", "방식과 결과"],
        ["방법 선택 기준",
         "등급이 규칙으로 결정되는 구조여서 예측 모형은 성립하지 않음. 구조를 드러내는 교차표와 조건 통제 비교를 주로 사용"],
        ["등급 결정 규칙",
         "모형이 아니라 교차표로 확인. 내용정보 최고값과 결정등급의 일치율 99.93%(134,462건), 2017년 5월 이후 94,516건은 예외 0건"],
        ["예측 모형 (검증용)",
         "내용정보로 등급을 예측하면 AUC 1.000. 성능이 아니라 항등식의 근거이므로 결과는 폐기"],
        ["상향 조정 분류 모형",
         "로지스틱 회귀 · 5겹 교차검증. 설명 변수는 내용정보 7항목·종별·연도. AUC 0.818, 설명력 0.170"],
        ["매체 간 비교",
         "보유 조합이 완전히 동일한 건만, 양측 50건 이상인 조합만 사용. 10개 조합 평균 차이 +15.1%p"],
    ],
    "closing": "예측이 목적이 아니라 구조 확인이 목적이므로, 성능 지표는 결론의 근거가 아니라 참고값으로만 사용",
    "notes": """[대본]
분석 방식을 정리하고 넘어가겠습니다.
이 프로젝트는 등급을 예측하는 모형을 만드는 것이 목적이 아닙니다. 앞에서 보신 대로 등급이 이미
규칙으로 정해지는 구조여서 예측 모형이 성립하지 않기 때문입니다.
그래서 교차표와 조건을 맞춘 비교를 주로 썼고, 모형은 두 번만 사용했습니다.
한 번은 검증용입니다. 내용정보로 등급을 예측했더니 AUC가 1.000이 나왔고, 이것을 성능이 아니라
항등식의 근거로 해석하고 결과는 폐기했습니다.
다른 하나는 상향 조정 여부를 설명하는 분류 모형입니다. 상향인지 아닌지를 맞히는 이진 분류라
로지스틱 회귀를 썼고, 5겹 교차검증으로 AUC 0.818, 설명력 0.170이 나왔습니다.
매체 비교는 모형 없이, 조건이 완전히 같은 건끼리만 비교하는 방식으로 했습니다.

[보충]
· AUC는 0.5가 아무 정보 없이 찍는 수준이고 1에 가까울수록 잘 맞힙니다. 0.818은 일정한 경향이
  있다는 뜻이고, 설명력 0.170은 설명되지 않는 부분이 더 크다는 뜻입니다. 개별 판단의 여지가 넓습니다.
· 로지스틱 회귀를 쓴 이유는 답이 '상향/비상향' 두 가지이고, 어떤 항목이 얼마나 영향을 주는지를
  숫자로 확인해야 했기 때문입니다.
· 5겹 교차검증은 데이터를 다섯 조각으로 나눠 네 조각으로 학습하고 한 조각으로 채점하기를
  다섯 번 반복하는 방식입니다. 한 번만 나눠서 생기는 우연을 줄입니다.

[예상 질문]
· 왜 딥러닝이나 복잡한 모형을 쓰지 않았나 → 목적이 예측이 아니라 구조 확인이었고, 규칙이 이미
  드러난 상태라 복잡한 모형을 쓸 이유가 없었습니다.
· 표본이 부족한 구간은 어떻게 했나 → 종별 100건, 조합 50건, 신청등급 200건 미만은 화면과
  리포트에 제시하지 않았습니다.""",
}

DASHBOARD = {
    "eyebrow": "6. 산출물",
    "title": "대시보드를 만든 이유",
    "rows": [
        ["구분", "내용"],
        ["만든 이유",
         "리포트는 고정된 조건의 결과만 제시. 조건을 바꿔도 결론이 유지되는지 직접 확인할 필요가 있었음"],
        ["확인할 수 있는 것",
         "등급 결정 규칙의 일치율, 등급을 올린 항목, 신청등급 대비 조정, 매체 간 판정 차이, 공개 데이터의 포괄 범위"],
        ["조건 변경",
         "기간·영상물 종별·게임물 플랫폼·성인물 제외를 바꾸면 수치와 해석 문장이 함께 재산출"],
        ["한계의 내장",
         "4단계 기준을 고르면 비교가 성립하지 않는다는 경고를 표시. 표본 미달 구간은 제시하지 않음"],
        ["개별 건 확인",
         "조건에 맞는 원자료를 화면에서 조회하고 전건을 CSV 로 내려받기 가능"],
    ],
    "closing": "리포트가 '무엇을 확인했는가'라면, 대시보드는 '조건을 바꿔도 그러한가'를 직접 확인하는 도구",
    "notes": """[대본]
산출물 가운데 대시보드를 왜 만들었는지 말씀드리겠습니다.
리포트는 제가 정해 둔 조건에서 나온 결과만 보여 줍니다. 그런데 이 분석의 결론에는 조건이 붙어
있습니다. 성인물을 뺐는지, 기간을 어디까지 봤는지, 이진화 기준을 몇 단계로 잡았는지에 따라
숫자가 달라집니다.
그래서 보는 사람이 조건을 직접 바꿔 가며 결론이 유지되는지 확인할 수 있어야 한다고 판단했습니다.
필터를 바꾸면 수치뿐 아니라 화면 위쪽의 해석 문장도 다시 계산됩니다.
한계도 화면에 넣었습니다. 이진화 기준을 4단계로 고르면 비교가 성립하지 않는다는 경고가 뜨고,
표본이 부족한 구간은 아예 표시하지 않습니다.
마지막 화면에서는 조건에 맞는 개별 건을 직접 확인하고 전건을 내려받을 수 있습니다.

[보충]
· 리포트·대시보드·발표자료가 모두 같은 계산부를 읽으므로 서로 다른 숫자가 나올 수 없습니다.
· 다음 장에서 실제 화면을 보여 드립니다.

[예상 질문]
· 누가 쓰는 것을 상정했나 → 등급분류 결과를 인용하려는 사람입니다. 어느 경로의 분류분인지,
  종별을 통제했는지에 따라 값이 달라지므로 그것을 직접 확인할 수 있게 만들었습니다.""",
}


def dup_slide(prs, index: int):
    """슬라이드를 통째로 복제한다. 서식이 그대로 따라온다."""
    src = prs.slides[index]
    dst = prs.slides.add_slide(src.slide_layout)
    for shape in list(dst.shapes):
        shape._element.getparent().remove(shape._element)
    for shape in src.shapes:
        dst.shapes._spTree.append(copy.deepcopy(shape._element))
    # 노트 자리틀도 원본에서 가져온다. 없으면 노트 글을 넣을 수 없다.
    if dst.notes_slide.notes_text_frame is None:
        ph = src.notes_slide.notes_placeholder
        dst.notes_slide.shapes._spTree.append(copy.deepcopy(ph._element))
    return dst


def move_slide(prs, from_idx: int, to_idx: int) -> None:
    lst = prs.slides._sldIdLst
    el = list(lst)[from_idx]
    lst.remove(el)
    lst.insert(to_idx, el)


def set_para(para, new: str) -> None:
    runs = list(para.runs)
    if not runs:
        return
    runs[0].text = new
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


def set_frame(tf, new: str) -> None:
    for i, para in enumerate(tf.paragraphs):
        if i == 0:
            set_para(para, new)
        else:
            for r in list(para.runs):
                r._r.getparent().remove(r._r)


def add_closing(prs, slide, text: str) -> None:
    """「폐기한 결론」 장의 맺음 문장 상자를 복제해 같은 서식으로 한 줄을 넣는다."""
    donor = None
    for s in prs.slides:
        for shape in s.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip().startswith("시사점:"):
                donor = shape
                break
        if donor is not None:
            break
    if donor is None:
        return
    el = copy.deepcopy(donor._element)
    slide.shapes._spTree.append(el)
    new_shape = slide.shapes[-1]
    new_shape.top = Inches(5.35)
    set_frame(new_shape.text_frame, text)


def fill(slide, spec: dict) -> None:
    for shape in slide.shapes:
        if shape.has_text_frame and shape.text_frame.text.strip():
            size = None
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.font.size:
                        size = run.font.size.pt
                        break
                if size:
                    break
            if size and size >= 24:
                set_frame(shape.text_frame, spec["title"])
            elif size and size <= 12:
                set_frame(shape.text_frame, spec["eyebrow"])
        elif shape.has_table:
            tbl = shape.table
            for ri, row in enumerate(spec["rows"]):
                for ci, val in enumerate(row):
                    set_frame(tbl.cell(ri, ci).text_frame, val)
    slide.notes_slide.notes_text_frame.text = spec["notes"].strip()


def main() -> None:
    prs = Presentation(DECK)
    titles = [next((sh.text_frame.text.strip() for sh in s.shapes
                    if sh.has_text_frame and sh.text_frame.text.strip()), "")
              for s in prs.slides]
    if METHOD["title"] in "".join(titles):
        print("[건너뜀] 이미 추가되어 있습니다.")
        return

    dash_idx = next(i for i, s in enumerate(prs.slides)
                    if any(sh.has_text_frame and DASHBOARD_TITLE in sh.text_frame.text
                           for sh in s.shapes))

    a = dup_slide(prs, TEMPLATE)
    fill(a, METHOD)
    add_closing(prs, a, METHOD["closing"])
    move_slide(prs, len(prs.slides._sldIdLst) - 1, 11)      # 결론 3 뒤 = 12번째

    dash_idx += 1                                            # 한 장 밀렸다
    b = dup_slide(prs, TEMPLATE + 1)                         # 위치가 밀린 원본
    fill(b, DASHBOARD)
    add_closing(prs, b, DASHBOARD["closing"])
    move_slide(prs, len(prs.slides._sldIdLst) - 1, dash_idx)

    prs.save(DECK)
    print(f"두 장 추가 완료 · 총 {len(prs.slides._sldIdLst)}장")


if __name__ == "__main__":
    main()
