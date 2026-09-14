# -*- coding: utf-8 -*-
"""
특강 제출용 작업기록(docx) 생성

WORKLOG.md 의 그날 항목은 상세본이라 그대로 두고, 제출용은 여기서 따로 만든다.
번호를 매긴 짧은 문장만 남기고, 근거·경위·시행착오는 넣지 않는다.

실행
    python docs/make_worklog.py
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from make_prd import (init_doc, h, para, bullet, table,
                      set_run_font, ACCENT, MUTED)                   # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH                        # noqa: E402

OUT_DIR = r"C:\Users\User\gucc_rating_analysis\docs"
DAY = "20260914"   # 이 값만 바꾸면 파일명과 제목이 함께 따라간다
OUT_FILE = os.path.join(OUT_DIR, "[BI 분석가 과정] 작업기록_%s_최은혜.docx" % DAY[2:])

DONE = [
    "통합 분석 Dataset 구축 — 게임물·영상물 240,290행 35열(2007~2026)",
    "내용정보 통합 규칙 확정 — 명칭이 동일한 4개 항목만 통합, 영상물 단계값은 원값 병기",
    "통합본 데이터 사전 작성 — 열 정의·통합 규칙·비교 가능 범위 명시",
    "청소년 이용 제한 콘텐츠 특성 분석 — 규모·추이·분포·항목·조합·자체등급분류 6개 절",
    "인사이트 리포트에 산출물 목록 절 추가 — 5종을 표로 정리하고 상호 연결",
    "대시보드 공개 배포 및 동작 확인 — 6개 화면",
    "PRD v1.7 개정 — 화면·필터·산출물·일정을 실제 구현에 맞춰 정정 및 문체 통일",
]

FOUND = [
    ["게임물", "청소년이용불가를 결정하는 요인은 폭력성이 아니라 사행성(단독 보유 97.1%)"],
    ["영상물", "청소년관람불가 43.7%의 대부분이 성인물이며, 제외 시 10.7%"],
    ["자체등급분류", "청소년이용불가 83건은 사업자 부여분이 아니라 사후 등급 조정의 결과"],
    ["매체 비교", "비교 가능한 10개 조합 전부에서 게임물의 제한 비율이 높음(평균 15.1%p, 3단계 기준)"],
    ["설계", "통합본은 제외 대상 건을 삭제하지 않고 표시만 함(제외 기준은 분석 목적에 따라 상이)"],
]

ISSUES = [
    "매체 비교 결론은 영상물의 3단계 절단 기준을 전제로 성립하며, 해당 기준은 데이터로 도출되지 않음",
    "무료 요금제 특성상 미사용 시 앱이 대기 상태로 전환되므로, 시연 전 사전 구동 필요",
    "통합본 CSV(59MB)는 저장소에서 제외하였으므로 스크립트로 재생성하거나 공유 폴더에서 확보 필요",
]


LINKS = [
    ["통합 분석 Dataset", "integrated_ratings_260914.parquet · .csv (240,290행)"],
    ["통합본 데이터 사전", "outputs/integrated_dataset_dictionary.html"],
    ["연령등급·내용정보 EDA", "outputs/eda_report.html · outputs/eda_grac_report.html"],
    ["청소년 이용 제한 특성 분석", "outputs/youth_report.html"],
    ["인터랙티브 대시보드", "https://gucc-rating-analysis-260910.streamlit.app/"],
    ["주요 분석 인사이트 리포트", "outputs/final_report.html"],
]


def build():
    doc = init_doc()

    p = doc.add_paragraph()
    r = p.add_run("%s - 최은혜" % DAY)
    set_run_font(r, size=16, bold=True, color=ACCENT)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = 14
    r = p.add_run("주제: 게임·영상물 등급분류 공공데이터 기반 "
                  "청소년 이용 부적합 콘텐츠 특성 분석 및 시각화")
    set_run_font(r, size=10, color=MUTED)

    h(doc, "오늘 한 일", 1)
    for i, t in enumerate(DONE, 1):
        para(doc, "%d. %s" % (i, t))

    h(doc, "오늘 나온 것", 1)
    table(doc, ["구분", "내용"], FOUND, widths=[2.6, 13.4])

    h(doc, "산출물", 1)
    table(doc, ["산출물", "위치"], LINKS, widths=[5.4, 10.6])

    h(doc, "이슈", 1)
    for t in ISSUES:
        bullet(doc, t)

    doc.save(OUT_FILE)
    print("saved:", OUT_FILE)


if __name__ == "__main__":
    build()
